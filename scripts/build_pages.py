"""키워드 원본 엑셀 → data/site-pages.json (사이트 라우트의 단일 출처) + keywords/대전철거_999.xlsx

  python -I scripts/build_pages.py

원본: source/keywords_src.xlsx (= 대전철거 1000개 키워드.xlsx) 첫 시트.
  A No · B 키워드 · C 붙여쓰기 · D 우선순위 · E 권역 · F 지역 표기 · G 지역 단위 · H 단어 · I 수식어
키워드 1개 = 페이지 1개, 파일 순서 유지.
오픈 단계(wave) = 우선순위: 1 = 대전 대표 18 / 2 = 전체 999
표기 변형(대전/대전시/대전광역시, 유성/유성구/대전 유성구/대전 유성, 둔산/둔산동, 도안/도안신도시)은
같은 장소·같은 단어·같은 수식어의 대표(우선순위↑, 파일 앞) 페이지를 variantOf 로.
"""
import json, sys
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))
from places_src import place_key  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'source' / 'keywords_src.xlsx'
OUT_XLSX = ROOT / 'keywords' / '대전철거_999.xlsx'
SITE_URL = 'https://daejeon-demo.example'  # TODO: 도메인 확정 후 lib/site.ts 와 함께 바꿀 것
UNIT = {'시·구': '시구', '생활권': '생활권'}

ws = load_workbook(SRC, read_only=True, data_only=True).worksheets[0]
rows = [r for r in list(ws.iter_rows(values_only=True))[1:] if r[1]]
pages = []
for i, r in enumerate(rows, 1):
    kw = str(r[1]).strip()
    pages.append(dict(n=i, no=r[0], kw=kw, slug=kw.replace(' ', '-'), label=r[5], place=place_key(r[5]), unit=UNIT[r[6]],
                      word=r[7], mod=r[8] or '', prio=int(r[3]), wave=int(r[3])))
assert len({p['slug'] for p in pages}) == len(pages), '슬러그 중복'

best = {}
for p in pages:
    k = (p['place'], p['word'], p['mod'])
    if k not in best or (p['prio'], p['n']) < (best[k]['prio'], best[k]['n']):
        best[k] = p
for p in pages:
    b = best[(p['place'], p['word'], p['mod'])]
    if b is not p:
        p['variantOf'] = b['slug']

(ROOT / 'data' / 'site-pages.json').write_text(json.dumps({'pages': pages}, ensure_ascii=False, indent=0), encoding='utf-8')

# 작업용 엑셀 (키워드 시트가 첫 시트)
wb = Workbook(); sh = wb.active; sh.title = '키워드'
sh.append(['번호', '원본 No', '키워드', '우선순위(단계)', '지역 표기', '장소', '지역 단위', '단어', '수식어', 'URL', '대표 페이지(표기 변형)'])
for p in pages:
    sh.append([p['n'], p['no'], p['kw'], p['wave'], p['label'], p['place'], p['unit'], p['word'], p['mod'],
               f"{SITE_URL}/{p['slug']}/", p.get('variantOf', '')])
for c in sh[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1A1D20'); c.alignment = Alignment(horizontal='center')
for col, w in zip('ABCDEFGHIJK', [6, 8, 28, 8, 14, 12, 8, 10, 8, 50, 28]):
    sh.column_dimensions[col].width = w
sh.freeze_panes = 'A2'
OUT_XLSX.parent.mkdir(exist_ok=True)
wb.save(OUT_XLSX)

reps = [p for p in pages if 'variantOf' not in p]
print('페이지', len(pages), '대표', len(reps), '변형', len(pages) - len(reps))
for w in (1, 2):
    op = [p for p in pages if p['wave'] <= w]
    print(f'{w}단계: 열린 {len(op)} / 사이트맵 키워드 {sum(1 for p in op if "variantOf" not in p or best[(p["place"], p["word"], p["mod"])]["wave"] > w)}')
print(Counter(p['word'] for p in pages))
